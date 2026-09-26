"""Preview or import a Zoho Students Pipeline workbook into the CRM.

The command is a dry run unless ``--apply`` is supplied. Matching is performed
by Zoho Record Id first. For records imported before that field existed, a
unique email or phone match is used once to attach the Zoho identifier.
"""

from __future__ import annotations

import argparse
import asyncio
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core.database import engine


STAGE_ALIASES = {
    "new inquiry": "new_inquiry",
    "new lead": "new_inquiry",
    "new leads": "new_inquiry",
    "potential lead": "new_inquiry",
    "contacted": "contacted",
    "attempted to contact": "contacted",
    "uncontactable no info": "contacted",
    "counselling": "counselling",
    "counseling": "counselling",
    "interested": "counselling",
    "follow up": "counselling",
    "exploring other options": "counselling",
    "waiting for parental approval": "counselling",
    "application started": "application_started",
    "future prospect": "application_started",
    "documents pending": "documents_pending",
    "looking for results": "documents_pending",
    "application submitted": "app_submitted",
    "app submitted": "app_submitted",
    "offer sent": "offer_sent",
    "ready to enrolled": "offer_sent",
    "payment done": "enrolled",
    "enrolled": "enrolled",
    "not interested declined": "lost_deferred",
    "not interested": "lost_deferred",
    "declined": "lost_deferred",
    "lost": "lost_deferred",
    "deferred": "lost_deferred",
    "junk leads": "lost_deferred",
    "can t afford": "lost_deferred",
    "lost to competitor": "lost_deferred",
}

FIELDS = (
    "full_name",
    "email",
    "phone",
    "whatsapp",
    "city",
    "highest_qualification",
    "interested_programme",
    "interested_course",
    "message",
    "source",
    "stage",
    "priority",
    "counsellor_name",
    "notes",
    "followup_date",
    "amount",
    "school",
    "nationality",
    "faculty",
    "student_status",
    "parents_occupation",
    "parents_email",
    "address_line1",
    "address_line2",
    "country",
    "social_lead_id",
    "external_record_id",
)


def clean(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = str(value).strip()
    return result or None


def normalize_words(value: Any) -> str:
    result = clean(value) or ""
    result = re.sub(r"[_/+-]+", " ", result.lower())
    return re.sub(r"[^a-z0-9]+", " ", result).strip()


def normalize_email(value: Any) -> str | None:
    result = (clean(value) or "").lower()
    return result if "@" in result else None


def normalize_phone(value: Any) -> str | None:
    digits = re.sub(r"\D", "", clean(value) or "")
    if len(digits) < 7:
        return None
    # Treat 07x, +947x, and 947x as the same Sri Lankan number.
    return digits[-9:] if len(digits) >= 9 else digits


def parse_amount(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return None


def parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    raw = clean(value)
    if not raw:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(raw[:19], fmt)
        except ValueError:
            continue
    return None


def mapped_stage(stage: Any, student_status: Any) -> tuple[str, bool]:
    raw = normalize_words(stage)
    if raw in STAGE_ALIASES:
        return STAGE_ALIASES[raw], True
    status = normalize_words(student_status)
    if status in STAGE_ALIASES:
        return STAGE_ALIASES[status], True
    if "payment" in status or "enrol" in status:
        return "enrolled", True
    if any(token in status for token in ("not interested", "declin", "defer")):
        return "lost_deferred", True
    return "new_inquiry", not bool(raw)


def cell(row: tuple[Any, ...], indexes: dict[str, int], name: str) -> Any:
    index = indexes.get(name)
    return row[index] if index is not None and index < len(row) else None


def read_workbook(path: Path) -> tuple[list[dict[str, Any]], Counter[str]]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    headers = [clean(item.value) for item in next(sheet.iter_rows())]
    indexes = {name: index for index, name in enumerate(headers) if name}
    required = {"Record Id", "Students Pipeline Name", "Stage"}
    missing = required.difference(indexes)
    if missing:
        raise ValueError(f"Workbook is missing required columns: {sorted(missing)}")

    records: list[dict[str, Any]] = []
    unknown_stages: Counter[str] = Counter()
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        full_name = clean(cell(row, indexes, "Students Pipeline Name")) or clean(
            cell(row, indexes, "Contact Name")
        )
        external_id = clean(cell(row, indexes, "Record Id"))
        if not full_name or not external_id:
            continue
        raw_stage = clean(cell(row, indexes, "Stage"))
        student_status = clean(cell(row, indexes, "Student Status"))
        stage, recognized = mapped_stage(raw_stage, student_status)
        if not recognized:
            unknown_stages[raw_stage or "(blank)"] += 1
        email = normalize_email(
            cell(row, indexes, "Email") or cell(row, indexes, "Email 1")
        )
        phone_raw = clean(cell(row, indexes, "Phone")) or clean(
            cell(row, indexes, "Phone 1")
        ) or clean(cell(row, indexes, "Mobile"))
        mobile_raw = clean(cell(row, indexes, "Mobile")) or phone_raw
        reason = clean(cell(row, indexes, "Reason For Loss"))
        source_raw = clean(cell(row, indexes, "Lead Source")) or "zoho_import"
        source = re.sub(r"[^a-z0-9]+", "_", source_raw.lower()).strip("_")
        priority = "high" if stage in {"offer_sent", "enrolled"} else "medium"
        records.append(
            {
                "_row": row_number,
                "_email_key": email,
                "_phone_key": normalize_phone(phone_raw),
                "full_name": full_name,
                "email": email,
                "phone": phone_raw,
                "whatsapp": mobile_raw,
                "city": clean(cell(row, indexes, "City")),
                "highest_qualification": clean(
                    cell(row, indexes, "Student Highest Education Qualification")
                ),
                "interested_programme": clean(cell(row, indexes, "Faculty  (Schools)")),
                "interested_course": clean(cell(row, indexes, "Selected Program")),
                "message": clean(cell(row, indexes, "Description")),
                "source": source or "zoho_import",
                "stage": stage,
                "priority": priority,
                "counsellor_name": clean(cell(row, indexes, "Students Pipeline Owner")),
                "notes": f"Reason for loss: {reason}" if reason else None,
                "followup_date": parse_datetime(cell(row, indexes, "Closing Date")),
                "amount": parse_amount(cell(row, indexes, "Amount")),
                "school": clean(cell(row, indexes, "School")),
                "nationality": clean(cell(row, indexes, "Nationality")),
                "faculty": clean(cell(row, indexes, "Faculty  (Schools)")),
                "student_status": student_status,
                "parents_occupation": clean(cell(row, indexes, "Parents Occupation")),
                "parents_email": normalize_email(cell(row, indexes, "Parents Email")),
                "address_line1": clean(cell(row, indexes, "Address Line 1")),
                "address_line2": clean(cell(row, indexes, "Address Line 2")),
                "country": clean(cell(row, indexes, "Country")),
                "social_lead_id": clean(cell(row, indexes, "Social Lead ID")),
                "external_record_id": external_id,
                "_created_at": parse_datetime(cell(row, indexes, "Created Time")),
                "_updated_at": parse_datetime(cell(row, indexes, "Modified Time")),
            }
        )
    return records, unknown_stages


async def column_exists(connection: Any) -> bool:
    result = await connection.scalar(
        text(
            "SELECT EXISTS (SELECT 1 FROM information_schema.columns "
            "WHERE table_schema='public' AND table_name='crm_leads' "
            "AND column_name='external_record_id')"
        )
    )
    return bool(result)


def build_lookup(rows: list[dict[str, Any]], field: str) -> dict[str, list[int]]:
    result: dict[str, list[int]] = defaultdict(list)
    for row in rows:
        value = row.get(field)
        if value:
            result[value].append(row["lead_id"])
    return result


async def run(path: Path, apply: bool) -> None:
    records, unknown_stages = read_workbook(path)
    external_counts = Counter(row["external_record_id"] for row in records)
    duplicate_external_ids = {key for key, count in external_counts.items() if count > 1}
    records = [row for row in records if row["external_record_id"] not in duplicate_external_ids]
    workbook_email_counts = Counter(row["_email_key"] for row in records if row["_email_key"])
    workbook_phone_counts = Counter(row["_phone_key"] for row in records if row["_phone_key"])

    async with engine.begin() as connection:
        has_external_id = await column_exists(connection)
        if apply and not has_external_id:
            raise RuntimeError(
                "Run scripts/apply_crm_external_record_id_migration.py before --apply."
            )
        external_select = "external_record_id" if has_external_id else "NULL AS external_record_id"
        existing_result = await connection.execute(
            text(
                "SELECT lead_id, lower(email) AS email_key, phone, "
                f"{external_select} FROM crm_leads"
            )
        )
        existing = []
        for item in existing_result.mappings():
            existing.append(
                {
                    "lead_id": item["lead_id"],
                    "email_key": normalize_email(item["email_key"]),
                    "phone_key": normalize_phone(item["phone"]),
                    "external_record_id": clean(item["external_record_id"]),
                }
            )

        by_external = {
            row["external_record_id"]: row["lead_id"]
            for row in existing
            if row["external_record_id"]
        }
        by_email = build_lookup(existing, "email_key")
        by_phone = build_lookup(existing, "phone_key")

        creates: list[dict[str, Any]] = []
        updates: list[dict[str, Any]] = []
        conflicts: list[tuple[int, str]] = []
        matched_external = matched_email = matched_phone = 0
        claimed_existing: set[int] = set()

        for record in records:
            lead_id = by_external.get(record["external_record_id"])
            match_kind = "external"
            if lead_id is None:
                candidates: set[int] = set()
                email_key = record["_email_key"]
                phone_key = record["_phone_key"]
                if email_key and workbook_email_counts[email_key] == 1:
                    email_matches = by_email.get(email_key, [])
                    if len(email_matches) == 1:
                        candidates.add(email_matches[0])
                if phone_key and workbook_phone_counts[phone_key] == 1:
                    phone_matches = by_phone.get(phone_key, [])
                    if len(phone_matches) == 1:
                        candidates.add(phone_matches[0])
                if len(candidates) > 1:
                    conflicts.append((record["_row"], "email and phone match different CRM records"))
                    continue
                if candidates:
                    lead_id = candidates.pop()
                    match_kind = "email" if email_key and lead_id in by_email.get(email_key, []) else "phone"
            values = {field: record.get(field) for field in FIELDS}
            if lead_id is None:
                values["created_at"] = record["_created_at"] or datetime.now()
                values["updated_at"] = record["_updated_at"] or values["created_at"]
                creates.append(values)
            elif lead_id in claimed_existing and match_kind != "external":
                conflicts.append((record["_row"], "fallback match was already claimed"))
            else:
                claimed_existing.add(lead_id)
                values["lead_id"] = lead_id
                values["import_updated_at"] = record["_updated_at"] or datetime.now()
                updates.append(values)
                if match_kind == "external":
                    matched_external += 1
                elif match_kind == "email":
                    matched_email += 1
                else:
                    matched_phone += 1

        print(f"Workbook rows ready: {len(records):,}")
        print(f"Existing CRM leads: {len(existing):,}")
        print(f"Would create: {len(creates):,}")
        print(f"Would update: {len(updates):,}")
        print(
            "Matches: "
            f"Zoho ID={matched_external:,}, email={matched_email:,}, phone={matched_phone:,}"
        )
        print(f"Conflicts skipped: {len(conflicts):,}")
        print(f"Duplicate Zoho IDs skipped: {len(duplicate_external_ids):,}")
        if unknown_stages:
            print(f"Unrecognized stage rows defaulted to new inquiry: {sum(unknown_stages.values()):,}")
            print("Unrecognized stages:", unknown_stages.most_common(20))
        if conflicts:
            print("First conflicts:", conflicts[:10])

        if not apply:
            print("Dry run only; no CRM records were changed.")
            return

        insert_fields = list(FIELDS) + ["created_at", "updated_at"]
        insert_columns = ", ".join(insert_fields)
        insert_values = ", ".join(f":{field}" for field in insert_fields)
        if creates:
            await connection.execute(
                text(f"INSERT INTO crm_leads ({insert_columns}) VALUES ({insert_values})"),
                creates,
            )

        assignments = ", ".join(
            f"{field} = COALESCE(:{field}, {field})" for field in FIELDS
        )
        if updates:
            await connection.execute(
                text(
                    f"UPDATE crm_leads SET {assignments}, "
                    "updated_at = :import_updated_at WHERE lead_id = :lead_id"
                ),
                updates,
            )
        print(f"Applied successfully: {len(creates):,} created, {len(updates):,} updated.")

    await engine.dispose()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    asyncio.run(run(args.workbook.resolve(), args.apply))


if __name__ == "__main__":
    main()
