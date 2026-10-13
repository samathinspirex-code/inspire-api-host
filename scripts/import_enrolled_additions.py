"""Review and import only the appended rows in Enrolled (1).xlsx.

The first 33 rows must match the previously reviewed workbook. Decisions about
the repeated student and Shenoli's account are required before --apply.
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.core.database import engine
from scripts.import_enrolled_workbook import HEADERS, clean, key, phone


SOURCE_SHA256 = "476c5818befee2f2e28b07ed3fb77506443884601c3e27f07537bd8b28759664"
SOURCE_NAME = "Enrolled (1).xlsx"
COURSE_NAMES = {
    "hnd in computing software engineering software developer pathway": "HND in Computing (Software Engineering)",
    "hnd in business management": "HND in Business Management",
}
APPROVED_MATCHES = {49: {24965}, 56: {19867}}


def read_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = iter(sheet.values)
    headers = tuple(clean(v) for v in next(values))
    if headers != HEADERS:
        raise ValueError("Workbook columns differ from the reviewed source")
    rows = [
        dict({header: cells[index] if index < len(cells) else None for index, header in enumerate(headers)}, _row=number)
        for number, cells in enumerate(values, 2) if any(v is not None for v in cells)
    ]
    if len(rows) != 60 or [r["_row"] for r in rows] != list(range(2, 62)):
        raise ValueError("Expected 60 contiguous data rows, ending at Excel row 61")
    return rows


def numbers(value: object) -> list[tuple[str, str]]:
    result = []
    for part in clean(value).split(","):
        normalized = phone(part)
        if normalized[0] and normalized not in result:
            result.append(normalized)
    return result


def same_original_rows(new_rows: list[dict], original: Path) -> None:
    old = load_workbook(original, read_only=True, data_only=True).active
    old_values = list(old.values)
    if tuple(clean(v) for v in old_values[0]) != HEADERS or len(old_values) != 34:
        raise ValueError("Original Enrolled.xlsx no longer has its 33 reviewed rows")
    for index, row in enumerate(new_rows[:33], start=1):
        if tuple(row[h] for h in HEADERS) != tuple(old_values[index]):
            raise ValueError(f"Existing Excel row {index + 1} changed; do not treat this as additions only")


def get_decisions(path: Path) -> dict:
    decisions = json.loads(path.read_text(encoding="utf-8"))
    if decisions.get("source_sha256") != SOURCE_SHA256:
        raise ValueError("Decisions are for a different workbook")
    if decisions.get("duplicate_rows_43_45") not in {"pending", "same_import_45", "different_import_both", "hold_both"}:
        raise ValueError("Invalid decision for rows 43 and 45")
    if decisions.get("shenoli_identity") not in {"pending", "existing_user_22", "new_account", "hold_rows"}:
        raise ValueError("Invalid Shenoli identity decision")
    return decisions


async def run(source: Path, original: Path, review: Path, apply: bool) -> None:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError("The workbook changed since review")
    rows = read_rows(source)
    same_original_rows(rows, original)
    decisions = get_decisions(review)
    additions = rows[33:]
    prepared = []
    for row in additions:
        row_number = row["_row"]
        if key(row["Status"]) != "payment done" or key(row["Lead Source"]) != "meta whatsapp":
            raise ValueError(f"Unexpected status or source on row {row_number}")
        if key(row["Created By"]) != "shenoli chandrawansa" or clean(row["Email"]):
            raise ValueError(f"Unexpected owner or email on row {row_number}")
        if not isinstance(row["Created Time"], datetime) or row["Created Time"].year != 2026:
            raise ValueError(f"Unexpected creation date on row {row_number}")
        if not isinstance(row["Closing Date"], datetime) or row["Closing Date"].year != 2026:
            raise ValueError(f"Unexpected closing date on row {row_number}")
        if clean(row["Amount"]) != "400000":
            raise ValueError(f"Unexpected amount on row {row_number}")
        target = COURSE_NAMES.get(key(row["Selected Program"]))
        if not target:
            raise ValueError(f"Unreviewed programme on row {row_number}: {row['Selected Program']!r}")
        parsed = numbers(row["Phone"]) + numbers(row["Mobile"])
        unique = list(dict(parsed).items())
        primary = unique[0][1] if unique else None
        secondary = unique[1][1] if len(unique) > 1 else None
        if len(unique) > 2:
            raise ValueError(f"More than two phone numbers on row {row_number}")
        prepared.append((row, target, unique, primary, secondary))

    by_phone: dict[str, list[int]] = defaultdict(list)
    for row, _, unique, _, _ in prepared:
        for number, _ in unique:
            by_phone[number].append(row["_row"])
    repeated = {number: matches for number, matches in by_phone.items() if len(matches) > 1}
    if len(repeated) != 1 or list(repeated.values()) != [[43, 45]]:
        raise ValueError(f"Unexpected repeated phone numbers in additions: {repeated}")

    async with engine.begin() as db:
        courses = (await db.execute(text("""
            SELECT c.course_id, c.title, c.awarding_body, p.name AS programme_name,
                   o.price AS programme_fee, o.duration
            FROM academic_courses c
            JOIN academic_programmes p ON p.programme_id=c.programme_id
            JOIN academic_schools s ON s.school_id=c.school_id
            LEFT JOIN LATERAL (
                SELECT price, duration FROM academic_course_study_options
                WHERE course_id=c.course_id AND is_enabled=true
                ORDER BY price, study_option_id LIMIT 1
            ) o ON true
            WHERE c.status='active' AND p.status='active' AND s.status='active'
        """))).mappings().all()
        catalogue = {key(course["title"]): dict(course) for course in courses}
        for target in set(COURSE_NAMES.values()):
            if key(target) not in catalogue:
                raise ValueError(f"Course no longer active: {target}")
        account = (await db.execute(text("""
            SELECT u.user_id, u.full_name, u.email, u.is_active,
                   coalesce(bool_or(al.access_key='COUNSELLOR'), false) AS is_counsellor
            FROM users u
            LEFT JOIN user_access_levels ua ON ua.user_id=u.user_id
            LEFT JOIN access_levels al ON al.access_level_id=ua.access_level_id
            WHERE u.user_id=22
            GROUP BY u.user_id, u.full_name, u.email, u.is_active
        """))).mappings().one_or_none()
        existing = (await db.execute(text("""
            SELECT lead_id, phone, whatsapp, external_record_id
            FROM crm_leads
        """))).mappings().all()
        existing_by_phone: dict[str, set[int]] = defaultdict(set)
        external_ids = set()
        for lead in existing:
            if lead["external_record_id"]:
                external_ids.add(lead["external_record_id"])
            for field in ("phone", "whatsapp"):
                normalized, _ = phone(lead[field])
                if normalized:
                    existing_by_phone[normalized].add(lead["lead_id"])
        matches = {}
        for row, _, unique, _, _ in prepared:
            row_number = row["_row"]
            external_id = f"enrolled-additions-{digest[:12]}-row-{row_number}"
            if external_id in external_ids:
                raise ValueError(f"Row {row_number} was already imported as {external_id}")
            ids = set().union(*(existing_by_phone[number] for number, _ in unique)) if unique else set()
            if ids != APPROVED_MATCHES.get(row_number, set()):
                raise ValueError(f"Unexpected CRM phone match on row {row_number}: {sorted(ids)}")
            if ids:
                matches[row_number] = sorted(ids)

        held_rows = []
        if decisions["duplicate_rows_43_45"] == "same_import_45":
            held_rows = [43]
        elif decisions["duplicate_rows_43_45"] == "hold_both":
            held_rows = [43, 45]
        if decisions["shenoli_identity"] == "hold_rows":
            held_rows = [row["_row"] for row in additions]
        ready = [item for item in prepared if item[0]["_row"] not in held_rows]
        print(json.dumps({
            "workbook_data_rows": len(rows), "previously_reviewed_rows": 33,
            "appended_rows": len(additions), "ready_rows": len(ready),
            "held_rows": held_rows, "approved_existing_matches": matches,
            "within_file_duplicate_rows": [43, 45],
            "blank_phone_rows": [r["_row"] for r, _, unique, _, _ in prepared if not unique],
            "second_number_rows": [r["_row"] for r, _, unique, _, _ in prepared if len(unique) > 1],
            "programme_counts": dict(Counter(target for _, target, _, _, _ in ready)),
            "existing_shenoli": dict(account) if account else None,
            "decisions": decisions,
        }, indent=2, default=str))
        if not apply:
            print("Preview only; no database changes.")
            return
        if "pending" in (decisions["duplicate_rows_43_45"], decisions["shenoli_identity"]):
            raise ValueError("Two review decisions are still pending")
        if not ready:
            print("No rows approved for import.")
            return
        if decisions["shenoli_identity"] == "existing_user_22":
            if not account or not account["is_active"] or key(account["full_name"]) != "shenoli":
                raise ValueError("Existing Shenoli account changed since review")
            owner_id, owner_name = account["user_id"], account["full_name"]
            if not account["is_counsellor"]:
                access_id = await db.scalar(text("SELECT access_level_id FROM access_levels WHERE access_key='COUNSELLOR' AND is_active=true"))
                if not access_id:
                    raise ValueError("COUNSELLOR access level unavailable")
                await db.execute(text("INSERT INTO user_access_levels (user_id, access_level_id) VALUES (:uid, :aid)"), {"uid": owner_id, "aid": access_id})
                if not await db.scalar(text("SELECT user_id FROM crm_counsellor_status WHERE user_id=:uid"), {"uid": owner_id}):
                    await db.execute(text("INSERT INTO crm_counsellor_status (user_id,is_active,assign_order) VALUES (:uid,false,9999)"), {"uid": owner_id})
        elif decisions["shenoli_identity"] == "new_account":
            if await db.scalar(text("SELECT user_id FROM users WHERE lower(full_name)='shenoli chandrawansa'")):
                raise ValueError("Shenoli Chandrawansa account already exists; review before retrying")
            placeholder = "shenoli.chandrawansa.legacy@example.com"
            if await db.scalar(text("SELECT user_id FROM users WHERE lower(email)=:email"), {"email": placeholder}):
                raise ValueError("Placeholder email already used")
            access_id = await db.scalar(text("SELECT access_level_id FROM access_levels WHERE access_key='COUNSELLOR' AND is_active=true"))
            if not access_id:
                raise ValueError("COUNSELLOR access level unavailable")
            owner_name = "Shenoli Chandrawansa"
            owner_id = await db.scalar(text("INSERT INTO users (full_name,email,is_active) VALUES (:name,:email,true) RETURNING user_id"), {"name": owner_name, "email": placeholder})
            await db.execute(text("INSERT INTO user_access_levels (user_id, access_level_id) VALUES (:uid,:aid)"), {"uid": owner_id, "aid": access_id})
            await db.execute(text("INSERT INTO crm_counsellor_status (user_id,is_active,assign_order) VALUES (:uid,false,9999)"), {"uid": owner_id})
        else:
            raise ValueError("No approved counsellor account")

        created = []
        for row, target, unique, primary, secondary in ready:
            row_number = row["_row"]
            course = catalogue[key(target)]
            audit = {field: clean(row[field]) for field in HEADERS}
            notes = f"Legacy import: {SOURCE_NAME}, row {row_number}. Original values: {json.dumps(audit, ensure_ascii=False)}"
            notes += "\nLegacy Enrolled exception: email pending; do not send offer or registration email until a valid address is supplied."
            notes += "\nWorkbook Amount is a down payment, not the full programme fee."
            if course["programme_fee"] is not None and 400000 > float(course["programme_fee"]):
                notes += "\nRecorded down payment exceeds current catalogue fee; verify amount and fee."
            if row_number in APPROVED_MATCHES:
                notes += f"\nDuplicate phone reviewed and approved as a separate lead; existing CRM lead(s) {matches[row_number]} unchanged."
            if row_number == 48:
                notes += "\nPhone matches an earlier held Digital Marketing workbook row; this lead was approved as a separate record."
            if row_number == 45 and decisions["duplicate_rows_43_45"] == "same_import_45":
                notes += "\nWorkbook row 43 is a duplicate of this student and was held by review."
            if not primary:
                notes += "\nWorkbook has no usable phone number; contact number pending."
            values = {
                "full_name": clean(row["Contact Name"]), "phone": primary,
                "whatsapp": secondary, "email": None, "source": "whatsapp", "stage": "enrolled",
                "priority": "medium", "interested_course": target,
                "interested_programme": course["programme_name"], "academic_course_id": course["course_id"],
                "awarding_body": course["awarding_body"], "programme_fee": course["programme_fee"],
                "programme_duration": course["duration"],
                "school": clean(row["School"]) or None, "nationality": clean(row["Nationality"]) or None,
                "faculty": clean(row["Faculty  (Schools)"]) or None,
                "highest_qualification": clean(row["Student Highest Education Qualification"]) or None,
                "city": clean(row["City"]) or None, "country": clean(row["Country"]) or None,
                "amount": 400000, "amount_paid": 400000, "payment_status": "down_payment",
                "payment_date": row["Closing Date"], "enrolled_at": row["Closing Date"],
                "created_at": row["Created Time"],
                "updated_at": row["Modified Time"] if isinstance(row["Modified Time"], datetime) else row["Created Time"],
                "assigned_at": row["Created Time"], "enrollment_email_status": "pending_email_legacy",
                "notes": notes, "external_record_id": f"enrolled-additions-{digest[:12]}-row-{row_number}",
                "is_archived": False, "assigned_counsellor_id": owner_id, "counsellor_name": owner_name,
            }
            columns = list(values)
            statement = text(f"INSERT INTO crm_leads ({', '.join(columns)}) VALUES ({', '.join(':'+c for c in columns)}) RETURNING lead_id")
            created.append((row_number, await db.scalar(statement, values)))
        print(json.dumps({"created_leads": created, "count": len(created), "held_rows": held_rows, "emails_sent": 0}, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--original", type=Path, required=True)
    parser.add_argument("--review", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    asyncio.run(run(args.source, args.original, args.review, args.apply))


if __name__ == "__main__":
    main()
