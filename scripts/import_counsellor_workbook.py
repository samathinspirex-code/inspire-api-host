"""Import the operational counsellor workbook into crm_leads.

Matches existing records by phone. Links assigned_counsellor_id to active
COUNSELLOR / CRM users by first name. Later sheets overwrite earlier ones.

Dry run by default; pass --apply to write.
"""

from __future__ import annotations

import argparse
import asyncio
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core.database import engine

SKIP_NAMES = {
    "number",
    "name",
    "programme",
    "program",
    "source",
    "update",
    "counselour",
    "counsellor",
    "counselor",
    "counselor name",
    "councelor name",
    "july",
    "athe",
    "cpd",
    "status",
    "date",
    "contact",
    "location",
    "lead source",
    "remarks",
}

STAGE_RANK = {
    "new_inquiry": 0,
    "contacted": 1,
    "counselling": 2,
    "application_started": 3,
    "documents_pending": 4,
    "app_submitted": 5,
    "offer_sent": 6,
    "enrolled": 7,
    "lost_deferred": 8,
}


def clean(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return None
    if isinstance(value, date):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = str(value).strip()
    return result or None


def normalize_words(value: Any) -> str:
    result = clean(value) or ""
    result = re.sub(r"[_/+-]+", " ", result.lower())
    return re.sub(r"[^a-z0-9]+", " ", result).strip()


def normalize_phone(value: Any) -> str | None:
    if isinstance(value, datetime):
        return None
    raw = value
    if isinstance(value, float) and value.is_integer():
        raw = int(value)
    digits = re.sub(r"\D", "", str(raw or ""))
    if len(digits) < 7:
        return None
    if len(digits) > 15:
        return None
    return digits[-9:] if len(digits) >= 9 else digits


def display_phone(phone_key: str) -> str:
    if len(phone_key) == 9 and phone_key.startswith("7"):
        return f"0{phone_key}"
    return phone_key


def looks_like_name(value: Any) -> bool:
    text = clean(value)
    if not text:
        return False
    key = normalize_words(text)
    if key in SKIP_NAMES or len(text) > 48:
        return False
    if re.fullmatch(r"[0-9+\s\-()]+", text):
        return False
    if any(
        token in key
        for token in (
            "told",
            "asked",
            "sent",
            "not interested",
            "interested",
            "answered",
            "declined",
            "whatsapp",
        )
    ):
        return False
    return bool(re.search(r"[a-zA-Z]", text))


def looks_like_phone_name(value: str | None) -> bool:
    if not value:
        return True
    lowered = value.strip().lower()
    if lowered in {"not provided", "unknown", "no name"}:
        return True
    return bool(re.fullmatch(r"[0-9+\s\-()]+", value.strip()))


def mapped_stage(update: Any, status: Any) -> str:
    raw = normalize_words(f"{status or ''} {update or ''}")
    if not raw:
        return "new_inquiry"
    lost_tokens = (
        "not interested",
        "not intrested",
        "declin",
        "junk",
        "wrong number",
        "invalid number",
        "not in use",
        "won t join",
        "wont join",
        "won t be able",
        "cant join",
        "can t join",
        "won t be joining",
        "lost to",
    )
    if any(token in raw for token in lost_tokens):
        return "lost_deferred"
    if any(token in raw for token in ("registered", "enrolled", "payment done", " paid")):
        return "enrolled"
    if any(token in raw for token in ("uncontactable", "no answer", "voicemail", "line busy", "phone off")):
        return "contacted"
    if any(token in raw for token in ("contacted", "shared", "sent details", "follow up", "will get back")):
        return "contacted"
    return "new_inquiry"


def merge_stage(existing: str | None, incoming: str) -> str:
    if not existing:
        return incoming
    if incoming == "lost_deferred":
        return incoming
    if existing in {"enrolled", "offer_sent", "app_submitted"} and incoming != "enrolled":
        return existing
    if STAGE_RANK.get(incoming, 0) >= STAGE_RANK.get(existing, 0):
        return incoming
    return existing


def merge_name(existing: str | None, incoming: str | None) -> str:
    if incoming and not looks_like_phone_name(incoming):
        if not existing or looks_like_phone_name(existing):
            return incoming
    return existing or incoming or "Unknown"


def source_key(value: Any) -> str:
    raw = clean(value) or "counsellor_workbook"
    key = re.sub(r"[^a-z0-9]+", "_", raw.lower()).strip("_")
    aliases = {
        "whatsapp": "whatsapp",
        "watsapp": "whatsapp",
        "whastapp": "whatsapp",
        "meta_whatsapp": "whatsapp",
        "lead_form": "lead_form",
        "leadform": "lead_form",
        "lead_forms": "lead_form",
        "meta_lead_forms": "lead_form",
    }
    return aliases.get(key, key or "counsellor_workbook")


def first_name_key(value: str | None) -> str | None:
    words = normalize_words(value)
    if not words:
        return None
    first = words.split()[0]
    aliases = {"jeniffer": "jennifer", "jenniffer": "jennifer"}
    return aliases.get(first, first)


def extract_phone_and_name(value: Any) -> tuple[str | None, str | None]:
    text = clean(value)
    phone = normalize_phone(value)
    name = None
    if text and not phone:
        phone = normalize_phone(text)
    if text and looks_like_name(re.sub(r"[0-9+\s\-()]+", " ", text)):
        leftover = re.sub(r"[0-9+\s\-()pP:]+", " ", text).strip()
        if looks_like_name(leftover):
            name = leftover
    return phone, name


def add_record(
    records: dict[str, dict[str, Any]],
    *,
    phone: Any,
    name: Any,
    programme: Any,
    source: Any,
    counsellor: Any,
    notes: Any,
    status: Any = None,
    city: Any = None,
    sheet: str,
) -> None:
    phone_key, extra_name = extract_phone_and_name(phone)
    if phone_key is None:
        return
    full_name = clean(name) if looks_like_name(name) else extra_name
    counsellor_name = clean(counsellor)
    if counsellor_name and normalize_words(counsellor_name) in SKIP_NAMES:
        counsellor_name = None
    if counsellor_name and re.search(r"\d", counsellor_name):
        counsellor_name = None
    note_parts = [part for part in (clean(notes), clean(status)) if part]
    incoming = {
        "phone_key": phone_key,
        "phone": display_phone(phone_key),
        "whatsapp": display_phone(phone_key),
        "full_name": full_name,
        "interested_course": clean(programme),
        "source": source_key(source),
        "counsellor_name": counsellor_name,
        "notes": " | ".join(note_parts) if note_parts else None,
        "stage": mapped_stage(notes, status),
        "city": clean(city) if looks_like_name(city) else None,
        "sheet": sheet,
    }
    records[phone_key] = incoming


def read_workbook(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    records: dict[str, dict[str, Any]] = {}

    sheet = workbook["527 - 79"]
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        if index == 0:
            continue
        add_record(
            records,
            phone=row[1] if len(row) > 1 else None,
            name=row[2] if len(row) > 2 else None,
            programme=row[3] if len(row) > 3 else None,
            source=row[4] if len(row) > 4 else None,
            counsellor=row[5] if len(row) > 5 else None,
            notes=row[6] if len(row) > 6 else None,
            status=row[7] if len(row) > 7 else None,
            sheet="527-79",
        )

    past = workbook["Past Counselors"]
    current_counsellor = "Yazeed"
    for index, row in enumerate(past.iter_rows(values_only=True), start=1):
        first = clean(row[0]) if row else None
        if first and normalize_words(first) in {"yazeed", "nevoli"}:
            current_counsellor = first.title()
            continue
        add_record(
            records,
            phone=row[1] if len(row) > 1 else None,
            name=row[2] if len(row) > 2 else None,
            programme=row[4] if len(row) > 4 else None,
            source=row[5] if len(row) > 5 else None,
            counsellor=current_counsellor,
            notes=" | ".join(part for part in (clean(row[6]) if len(row) > 6 else None, clean(row[7]) if len(row) > 7 else None) if part),
            sheet="past_counselors",
        )

    samindi = workbook["Samindi"]
    for index, row in enumerate(samindi.iter_rows(values_only=True)):
        if index < 3:
            continue
        for phone_i, name_i, programme_i, source_i, notes_i in (
            (1, 2, 3, 4, 5),
            (11, 12, 13, 14, 15),
            (19, 20, 21, 22, 24),
        ):
            add_record(
                records,
                phone=row[phone_i] if len(row) > phone_i else None,
                name=row[name_i] if len(row) > name_i else None,
                programme=row[programme_i] if len(row) > programme_i else None,
                source=row[source_i] if len(row) > source_i else None,
                counsellor="Samindi",
                notes=row[notes_i] if len(row) > notes_i else None,
                sheet="samindi",
            )

    jennifer = workbook["Jennifer"]
    for index, row in enumerate(jennifer.iter_rows(values_only=True)):
        if index == 0:
            continue
        add_record(
            records,
            phone=row[1] if len(row) > 1 else None,
            name=row[2] if len(row) > 2 else None,
            programme=row[3] if len(row) > 3 else None,
            source=row[4] if len(row) > 4 else None,
            counsellor="Jennifer",
            notes=row[5] if len(row) > 5 else None,
            sheet="jennifer",
        )
        add_record(
            records,
            phone=row[9] if len(row) > 9 else None,
            name=row[10] if len(row) > 10 else None,
            programme=row[12] if len(row) > 12 else None,
            source=row[13] if len(row) > 13 else None,
            counsellor="Jennifer",
            notes=row[14] if len(row) > 14 else None,
            city=row[11] if len(row) > 11 else None,
            sheet="jennifer",
        )
        add_record(
            records,
            phone=row[17] if len(row) > 17 else None,
            name=row[18] if len(row) > 18 else None,
            programme=row[20] if len(row) > 20 else None,
            source=row[21] if len(row) > 21 else None,
            counsellor="Jennifer",
            notes=row[22] if len(row) > 22 else None,
            sheet="jennifer",
        )

    newest = workbook["New Counselors"]
    for index, row in enumerate(newest.iter_rows(values_only=True)):
        if index < 2:
            continue
        add_record(
            records,
            phone=row[3] if len(row) > 3 else None,
            name=row[4] if len(row) > 4 else None,
            programme=row[6] if len(row) > 6 else None,
            source=row[7] if len(row) > 7 else None,
            counsellor=row[5] if len(row) > 5 else None,
            notes=row[8] if len(row) > 8 else None,
            sheet="new_counselors",
        )
        add_record(
            records,
            phone=row[13] if len(row) > 13 else None,
            name=row[12] if len(row) > 12 else None,
            programme=row[15] if len(row) > 15 else None,
            source=row[16] if len(row) > 16 else None,
            counsellor=row[14] if len(row) > 14 else None,
            notes=row[17] if len(row) > 17 else None,
            sheet="new_counselors",
        )
        add_record(
            records,
            phone=row[22] if len(row) > 22 else None,
            name=row[21] if len(row) > 21 else None,
            programme=row[24] if len(row) > 24 else None,
            source=row[25] if len(row) > 25 else None,
            counsellor=row[23] if len(row) > 23 else None,
            notes=row[26] if len(row) > 26 else None,
            sheet="new_counselors",
        )

    return records


async def load_counsellor_index(connection: Any) -> dict[str, dict[str, Any]]:
    rows = (
        await connection.execute(
            text(
                """
                SELECT u.user_id, u.full_name, u.email,
                       bool_or(al.access_key = 'COUNSELLOR') AS is_counsellor,
                       bool_or(al.access_key = 'CRM') AS is_crm
                FROM users u
                JOIN user_access_levels ual ON ual.user_id = u.user_id
                JOIN access_levels al ON al.access_level_id = ual.access_level_id
                WHERE u.is_active = true
                  AND al.is_active = true
                  AND al.access_key IN ('COUNSELLOR', 'CRM')
                GROUP BY u.user_id, u.full_name, u.email
                ORDER BY bool_or(al.access_key = 'COUNSELLOR') DESC, u.user_id
                """
            )
        )
    ).mappings().all()
    index: dict[str, dict[str, Any]] = {}
    for row in rows:
        key = first_name_key(row["full_name"])
        if not key:
            continue
        current = index.get(key)
        if current is None or (row["is_counsellor"] and not current["is_counsellor"]):
            index[key] = dict(row)
    return index


def resolve_counsellor(
    name: str | None, index: dict[str, dict[str, Any]]
) -> tuple[int | None, str | None]:
    key = first_name_key(name)
    if not key:
        return None, name
    match = index.get(key)
    if match:
        return match["user_id"], match["full_name"]
    return None, name


async def run(path: Path, apply: bool) -> None:
    workbook_records = read_workbook(path)
    async with engine.begin() as connection:
        counsellor_index = await load_counsellor_index(connection)
        print("Counsellor users:")
        for key, row in counsellor_index.items():
            kind = "COUNSELLOR" if row["is_counsellor"] else "CRM"
            print(f"  {row['full_name']} #{row['user_id']} ({kind}) key={key}")

        existing_rows = (
            await connection.execute(
                text(
                    "SELECT lead_id, full_name, phone, email, stage, source, "
                    "interested_course, notes, counsellor_name, assigned_counsellor_id "
                    "FROM crm_leads"
                )
            )
        ).mappings().all()
        existing = [dict(row) for row in existing_rows]
        by_phone: dict[str, list[dict[str, Any]]] = {}
        for row in existing:
            key = normalize_phone(row["phone"])
            row["phone_key"] = key
            if key:
                by_phone.setdefault(key, []).append(row)

        link_only: list[dict[str, Any]] = []
        updates: list[dict[str, Any]] = []
        creates: list[dict[str, Any]] = []
        unmatched_counsellors: Counter[str] = Counter()
        claimed: set[int] = set()

        for row in existing:
            user_id, official_name = resolve_counsellor(row["counsellor_name"], counsellor_index)
            if user_id and row["assigned_counsellor_id"] != user_id:
                link_only.append(
                    {
                        "lead_id": row["lead_id"],
                        "assigned_counsellor_id": user_id,
                        "counsellor_name": official_name,
                    }
                )
                row["assigned_counsellor_id"] = user_id
                row["counsellor_name"] = official_name

        for record in workbook_records.values():
            user_id, official_name = resolve_counsellor(record["counsellor_name"], counsellor_index)
            if record["counsellor_name"] and user_id is None:
                unmatched_counsellors[record["counsellor_name"]] += 1
            record["assigned_counsellor_id"] = user_id
            if official_name:
                record["counsellor_name"] = official_name
            matches = [item for item in by_phone.get(record["phone_key"], []) if item["lead_id"] not in claimed]
            if matches:
                target = matches[0]
                claimed.add(target["lead_id"])
                incoming_name = record["counsellor_name"]
                incoming_id = record["assigned_counsellor_id"]
                if incoming_name and incoming_id is None:
                    assigned_name, assigned_id = incoming_name, None
                else:
                    assigned_name = incoming_name or target["counsellor_name"]
                    assigned_id = incoming_id or target["assigned_counsellor_id"]
                updates.append(
                    {
                        "lead_id": target["lead_id"],
                        "full_name": merge_name(target["full_name"], record["full_name"]),
                        "phone": record["phone"] or target["phone"],
                        "whatsapp": record["whatsapp"],
                        "interested_course": record["interested_course"] or target["interested_course"],
                        "source": record["source"] or target["source"],
                        "notes": record["notes"] or target["notes"],
                        "stage": merge_stage(target["stage"], record["stage"]),
                        "city": record["city"],
                        "counsellor_name": assigned_name,
                        "assigned_counsellor_id": assigned_id,
                        "priority": "high" if record["stage"] == "enrolled" else "medium",
                    }
                )
            else:
                creates.append(
                    {
                        "full_name": record["full_name"] or display_phone(record["phone_key"]),
                        "email": None,
                        "phone": record["phone"],
                        "whatsapp": record["whatsapp"],
                        "city": record["city"],
                        "interested_course": record["interested_course"],
                        "message": None,
                        "source": record["source"],
                        "stage": record["stage"],
                        "priority": "high" if record["stage"] == "enrolled" else "medium",
                        "assigned_counsellor_id": record["assigned_counsellor_id"],
                        "counsellor_name": record["counsellor_name"],
                        "notes": record["notes"],
                    }
                )

        print(f"Workbook unique phones: {len(workbook_records):,}")
        print(f"Existing CRM leads: {len(existing):,}")
        print(f"Would link counsellor IDs on existing leads: {len(link_only):,}")
        print(f"Would update from workbook: {len(updates):,}")
        print(f"Would create: {len(creates):,}")
        if unmatched_counsellors:
            print("Unmatched workbook counsellors:", unmatched_counsellors.most_common())

        assigned_creates = sum(1 for row in creates if row["assigned_counsellor_id"])
        assigned_updates = sum(1 for row in updates if row["assigned_counsellor_id"])
        print(f"Creates with counsellor user: {assigned_creates:,}")
        print(f"Workbook updates with counsellor user: {assigned_updates:,}")

        if not apply:
            print("Dry run only; no CRM records were changed.")
            return

        if link_only:
            await connection.execute(
                text(
                    "UPDATE crm_leads SET assigned_counsellor_id = :assigned_counsellor_id, "
                    "counsellor_name = :counsellor_name, updated_at = NOW() "
                    "WHERE lead_id = :lead_id"
                ),
                link_only,
            )

        if updates:
            await connection.execute(
                text(
                    """
                    UPDATE crm_leads SET
                        full_name = :full_name,
                        phone = COALESCE(:phone, phone),
                        whatsapp = COALESCE(:whatsapp, whatsapp),
                        interested_course = COALESCE(:interested_course, interested_course),
                        source = COALESCE(:source, source),
                        notes = COALESCE(:notes, notes),
                        stage = :stage,
                        city = COALESCE(:city, city),
                        counsellor_name = :counsellor_name,
                        assigned_counsellor_id = :assigned_counsellor_id,
                        priority = :priority,
                        updated_at = NOW()
                    WHERE lead_id = :lead_id
                    """
                ),
                updates,
            )

        if creates:
            await connection.execute(
                text(
                    """
                    INSERT INTO crm_leads (
                        full_name, email, phone, whatsapp, city, interested_course,
                        message, source, stage, priority, assigned_counsellor_id,
                        counsellor_name, notes
                    ) VALUES (
                        :full_name, :email, :phone, :whatsapp, :city, :interested_course,
                        :message, :source, :stage, :priority, :assigned_counsellor_id,
                        :counsellor_name, :notes
                    )
                    """
                ),
                creates,
            )

        remaining_unassigned = int(
            await connection.scalar(
                text(
                    "SELECT COUNT(*) FROM crm_leads "
                    "WHERE assigned_counsellor_id IS NULL AND is_archived = false"
                )
            )
            or 0
        )
        print(
            f"Applied: {len(link_only):,} counsellor links, "
            f"{len(updates):,} workbook updates, {len(creates):,} created."
        )
        print(f"Leads still without a counsellor user: {remaining_unassigned:,}")

    await engine.dispose()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    asyncio.run(run(args.workbook.resolve(), args.apply))


if __name__ == "__main__":
    main()
