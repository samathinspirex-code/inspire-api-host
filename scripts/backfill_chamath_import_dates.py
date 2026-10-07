"""Backfill historical timestamps for leads imported from chamath.xlsx.

Run with a workbook path to preview; add --apply to update the database.
Only leads carrying chamath-row-* or chamath-merge-* external IDs are touched.
"""

import argparse
import asyncio
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.core.database import engine  # noqa: E402


EXTERNAL_ID = re.compile(r"^chamath-(?:row-(\d+)|merge-(\d+)-(\d+))$")


def workbook_dates(path: Path) -> dict[int, tuple[datetime, str]]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    if tuple(cell.value for cell in sheet[1])[:9] != (
        "Contact Name", "Created Time", "Lead Source", "Selected Program",
        "Status", "Phone", "Mobile", "Closing Date", "Created By",
    ):
        raise ValueError("Unexpected workbook columns; no dates were updated")
    values = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        raw = row[1]
        created = raw if isinstance(raw, datetime) else datetime.fromisoformat(str(raw).strip())
        values[row_number] = (created.replace(tzinfo=None), str(row[7] or ""))
    return values


async def run(path: Path, apply: bool) -> None:
    dates = workbook_dates(path)
    engine.echo = False
    async with engine.begin() as connection:
        result = await connection.execute(text(
            "SELECT lead_id, external_record_id, stage, created_at, assigned_at, "
            "updated_at, enrolled_at, notes FROM crm_leads "
            "WHERE external_record_id LIKE 'chamath-row-%' "
            "OR external_record_id LIKE 'chamath-merge-%'"
        ))
        leads = result.mappings().all()
        if len(leads) != 473:
            raise ValueError(f"Expected 473 imported leads, found {len(leads)}; no dates were updated")
        updates = []
        changes = []
        covered = set()
        for lead in leads:
            match = EXTERNAL_ID.fullmatch(lead["external_record_id"])
            if not match:
                raise ValueError(f"Unexpected import ID: {lead['external_record_id']}")
            row_numbers = [int(value) for value in match.groups() if value]
            if any(number not in dates for number in row_numbers):
                raise ValueError(f"Workbook row missing for {lead['external_record_id']}")
            covered.update(row_numbers)
            created = min(dates[number][0] for number in row_numbers)
            original_dates = "; ".join(
                f"row {number}: Created Time={dates[number][0].isoformat(sep=' ')}; Closing Date={dates[number][1]}"
                for number in row_numbers
            )
            notes = lead["notes"] or ""
            if "Original workbook dates:" not in notes:
                notes = f"{notes}\nOriginal workbook dates: {original_dates}".strip()
            update = {
                "id": lead["lead_id"],
                "created": created,
                "enrolled": created if lead["stage"] == "enrolled" else lead["enrolled_at"],
                "notes": notes,
            }
            updates.append(update)
            if lead["created_at"] != created or lead["assigned_at"] != created:
                changes.append(update)
        if len(covered) != 485:
            raise ValueError(f"Expected 485 represented workbook rows, found {len(covered)}")
        print(f"Imported leads: {len(leads)}; represented workbook rows: {len(covered)}")
        print(f"Historical Created Time range: {min(item['created'] for item in updates)} to {max(item['created'] for item in updates)}")
        print(f"Records needing date changes: {len(changes)}")
        print(f"Legacy Enrolled records: {sum(lead['stage'] == 'enrolled' for lead in leads)}")
        if apply:
            if changes:
                await connection.execute(text(
                    "UPDATE crm_leads SET created_at=:created, assigned_at=:created, "
                    "updated_at=:created, enrolled_at=:enrolled, notes=:notes WHERE lead_id=:id"
                ), changes)
                print("Historical dates applied atomically.")
            else:
                print("No date changes needed.")
        else:
            print("Preview only. Use --apply to commit these changes.")
    await engine.dispose()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--apply", action="store_true")
    arguments = parser.parse_args()
    asyncio.run(run(arguments.workbook, arguments.apply))
